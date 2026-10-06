"""
Top 5 de partidas de la balanza de pagos (BanRep, máximo nivel de detalle),
por cuenta: participación, crecimiento y contribución al crecimiento.

Para cada cuenta y cada lado
  * Cuenta corriente : neto, crédito (ingresos), débito (egresos)
  * Cuenta financiera: neto, activos (adquisición neta), pasivos (netos incurridos)
se calculan, sobre las partidas de máximo detalle (ver src/build/bp_banrep_detalle.py):

  participación (%)       = x_t / X_t * 100
  participación abs. (%)  = |x_t| / sum|x_t| * 100   (útil cuando hay signos mixtos)
  crecimiento (%)         = (x_t - x_b) / |x_b| * 100
  contribución (p.p.)     = (x_t - x_b) / |X_b| * 100  -> suman el crecimiento de X

donde t es el periodo analizado, b el periodo base y X el total del lado.

Rankings
  1. Top participación  : mayor participación absoluta en t (+ su crecimiento).
  2. Top crecimiento    : mayor crecimiento entre partidas con participación
                          absoluta >= --umbral (% en t o en b) (+ su participación).
  3. Top contribución   : mayor |contribución| (+ participación y crecimiento).

Uso
    python src/analisis/top_partidas_bp.py                    # último semestre vs mismo semestre año anterior
    python src/analisis/top_partidas_bp.py --periodo 2026S1
    python src/analisis/top_partidas_bp.py --periodo 2026Q2 --base anterior
    python src/analisis/top_partidas_bp.py --periodo 2025 --top 10 --umbral 0.5

Salida: output/top_partidas_bp_<periodo>_vs_<base>.xlsx
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src" / "build"))
import bp_banrep_detalle as bp  # noqa: E402

OUTPUT = ROOT / "output"

LADOS = {
    "1": [("neto", "Saldo"), ("credito", "Crédito (ingresos)"), ("debito", "Débito (egresos)")],
    "3": [("neto", "Neto"), ("activos", "Adquisición neta de activos"),
          ("pasivos", "Pasivos netos incurridos")],
}
SIGLA = {"1": "CC", "3": "CF"}


# ----------------------------------------------------------------------------
# Periodos
# ----------------------------------------------------------------------------
def trimestres(periodo: str) -> list[str]:
    """'2026S1' -> ['2026Q1','2026Q2'];  '2026Q2' -> ['2026Q2'];  '2025' -> Q1..Q4."""
    p = periodo.upper()
    if m := re.fullmatch(r"(\d{4})S([12])", p):
        a, s = int(m[1]), int(m[2])
        return [f"{a}Q{q}" for q in ((1, 2) if s == 1 else (3, 4))]
    if m := re.fullmatch(r"(\d{4})[QT]([1-4])", p):
        return [f"{m[1]}Q{m[2]}"]
    if re.fullmatch(r"\d{4}", p):
        return [f"{p}Q{q}" for q in range(1, 5)]
    raise ValueError(f"Periodo no reconocido: {periodo} (use 2026S1, 2026Q2 o 2025)")


def periodo_base(periodo: str, base: str) -> str:
    p = periodo.upper()
    if base == "anual":
        return str(int(p[:4]) - 1) + p[4:]
    if m := re.fullmatch(r"(\d{4})S([12])", p):
        a, s = int(m[1]), int(m[2])
        return f"{a - 1}S2" if s == 1 else f"{a}S1"
    if m := re.fullmatch(r"(\d{4})[QT]([1-4])", p):
        a, q = int(m[1]), int(m[2])
        return f"{a - 1}Q4" if q == 1 else f"{a}Q{q - 1}"
    return str(int(p) - 1)


def ultimo_semestre(ancho: pd.DataFrame) -> str:
    ult = ancho.columns[ancho.notna().any()].max()  # p. ej. '2026Q2'
    a, q = int(ult[:4]), int(ult[-1])
    if q >= 4:
        return f"{a}S2"
    if q >= 2:
        return f"{a}S1"
    return f"{a - 1}S2"


# ----------------------------------------------------------------------------
# Cálculo
# ----------------------------------------------------------------------------
def valores_lado(cuenta: str, lado: str, hojas: pd.DataFrame, ancho: pd.DataFrame,
                 lados_cc: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """Orden x trimestre con los valores de las hojas para ese lado."""
    h = hojas[hojas.cuenta == cuenta].set_index("Orden")
    if cuenta == "1" and lado in ("credito", "debito"):
        return lados_cc[lado].reindex(h.index)
    if lado == "neto":
        return ancho.reindex(h.index).mul(h["signo"], axis=0)
    return ancho.reindex(h.loc[h.lado == lado].index)


def metricas(x_t: pd.Series, x_b: pd.Series) -> pd.DataFrame:
    X_t, X_b = x_t.sum(), x_b.sum()
    df = pd.DataFrame({"valor_t": x_t, "valor_b": x_b})
    df["var_abs"] = df.valor_t - df.valor_b
    df["participacion_t"] = df.valor_t / X_t * 100 if X_t else np.nan
    df["participacion_b"] = df.valor_b / X_b * 100 if X_b else np.nan
    df["part_abs_t"] = df.valor_t.abs() / df.valor_t.abs().sum() * 100
    df["part_abs_b"] = df.valor_b.abs() / df.valor_b.abs().sum() * 100
    df["crecimiento"] = np.where(df.valor_b != 0, df.var_abs / df.valor_b.abs() * 100, np.nan)
    df["contribucion_pp"] = df.var_abs / abs(X_b) * 100 if X_b else np.nan
    return df


def rankings(m: pd.DataFrame, top: int, umbral: float) -> dict[str, pd.DataFrame]:
    m = m[(m.valor_t.fillna(0) != 0) | (m.valor_b.fillna(0) != 0)]
    material = m[((m.part_abs_t >= umbral) | (m.part_abs_b >= umbral)) & m.crecimiento.notna()]
    return {
        "participacion": m.sort_values("part_abs_t", ascending=False).head(top),
        "crecimiento": material.sort_values("crecimiento", ascending=False).head(top),
        "contribucion": m.reindex(m.contribucion_pp.abs().sort_values(ascending=False).index).head(top),
    }


COLS = {
    "participacion": ["participacion_t", "part_abs_t", "crecimiento", "valor_t", "valor_b", "var_abs"],
    "crecimiento": ["crecimiento", "participacion_t", "part_abs_t", "valor_t", "valor_b", "var_abs"],
    "contribucion": ["contribucion_pp", "participacion_t", "crecimiento", "valor_t", "valor_b", "var_abs"],
}
TITULOS = {
    "participacion": "Top {n} por participación en la cuenta",
    "crecimiento": "Top {n} por crecimiento (partidas con participación abs. ≥ {u}%)",
    "contribucion": "Top {n} por contribución al crecimiento de la cuenta",
}


def encabezados(t: str, b: str) -> dict[str, str]:
    return {
        "etiqueta": "Partida", "codigo": "Código BanRep",
        "valor_t": f"Valor {t} (USD mill.)", "valor_b": f"Valor {b} (USD mill.)",
        "var_abs": "Variación (USD mill.)", "participacion_t": f"Participación {t} (%)",
        "part_abs_t": f"Participación abs. {t} (%)", "crecimiento": "Crecimiento (%)",
        "contribucion_pp": "Contribución (p.p.)",
    }


# ----------------------------------------------------------------------------
# Excel
# ----------------------------------------------------------------------------
def escribir_excel(ruta: Path, bloques: list[dict], detalle: pd.DataFrame, control: pd.DataFrame,
                   notas: list[str], t: str, b: str) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    enc = encabezados(t, b)
    negrita, titulo = Font(bold=True), Font(bold=True, size=13)
    relleno = PatternFill("solid", fgColor="DDE6F0")
    with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
        pd.DataFrame({"Notas": notas}).to_excel(xw, sheet_name="Notas", index=False)
        for blq in bloques:
            hoja = blq["hoja"]
            fila = 3
            for clave, df in blq["rankings"].items():
                tabla = df.reset_index()[["etiqueta", "codigo"] + COLS[clave]].rename(columns=enc)
                tabla.insert(0, "#", range(1, len(tabla) + 1))
                pd.DataFrame([[TITULOS[clave].format(n=blq["top"], u=blq["umbral"])]]).to_excel(
                    xw, sheet_name=hoja, startrow=fila, index=False, header=False)
                tabla.to_excel(xw, sheet_name=hoja, startrow=fila + 1, index=False)
                ws = xw.sheets[hoja]
                ws.cell(fila + 1, 1).font = negrita
                for c in range(1, tabla.shape[1] + 1):
                    cel = ws.cell(fila + 2, c)
                    cel.font, cel.fill = negrita, relleno
                    cel.alignment = Alignment(wrap_text=True, vertical="center")
                for r in range(fila + 3, fila + 3 + len(tabla)):
                    for c in range(4, tabla.shape[1] + 1):
                        ws.cell(r, c).number_format = "#,##0.0"
                fila += len(tabla) + 4
            ws = xw.sheets[hoja]
            ws.cell(1, 1, blq["titulo"]).font = titulo
            ws.cell(2, 1, blq["subtitulo"])
            ws.column_dimensions["A"].width = 4
            ws.column_dimensions["B"].width = 85
            ws.column_dimensions["C"].width = 14
            for c in range(4, 10):
                ws.column_dimensions[get_column_letter(c)].width = 16

        detalle.to_excel(xw, sheet_name="Detalle_hojas", index=False)
        control.to_excel(xw, sheet_name="Control", index=False)
        for nombre in ("Detalle_hojas", "Control", "Notas"):
            ws = xw.sheets[nombre]
            ws.column_dimensions["A"].width = 110 if nombre == "Notas" else 22
            for c in ws[1]:
                c.font, c.fill = negrita, relleno
        xw.sheets["Detalle_hojas"].column_dimensions["D"].width = 85


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------
def main(argv: list[str] | None = None) -> Path:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--periodo", help="2026S1 (semestre), 2026Q2 (trimestre) o 2025 (año). "
                                      "Por defecto: último semestre completo.")
    ap.add_argument("--base", choices=["anual", "anterior"], default="anual",
                    help="anual = mismo periodo del año anterior (defecto); anterior = periodo inmediatamente anterior")
    ap.add_argument("--top", type=int, default=5)
    ap.add_argument("--umbral", type=float, default=1.0,
                    help="Participación abs. mínima (%%) para entrar al ranking de crecimiento (defecto 1)")
    args = ap.parse_args(argv)

    df = bp.leer_raw()
    arbol = bp.construir_arbol(df)
    largo = bp.cargar_largo(df)
    ancho = bp.a_ancho(largo)

    t = (args.periodo or ultimo_semestre(ancho)).upper()
    b = periodo_base(t, args.base)
    qt, qb = trimestres(t), trimestres(b)
    faltan = [q for q in qt + qb if q not in ancho.columns]
    if faltan:
        raise SystemExit(f"No hay datos para: {faltan}")

    # Hojas que cuadran en todos los trimestres comparados.
    hojas = bp.seleccionar_hojas(arbol, ancho, periodos=qb + qt, cuentas=("1", "3"))
    lados_cc = bp.lados_cuenta_corriente(arbol, ancho, hojas)
    info = hojas.set_index("Orden")[["etiqueta", "codigo", "lado", "cuenta"]]
    estados = largo[largo.periodo.isin(qt)].estado.dropna().unique()

    bloques, detalle, control = [], [], []
    for cuenta in ("1", "3"):
        nombre = bp.CUENTAS[cuenta]
        orden_total = arbol.index[(arbol.codigo == cuenta) & (arbol.nivel == 0)][0]
        for lado, nombre_lado in LADOS[cuenta]:
            v = valores_lado(cuenta, lado, hojas, ancho, lados_cc)
            x_t, x_b = v[qt].fillna(0).sum(axis=1), v[qb].fillna(0).sum(axis=1)
            m = metricas(x_t, x_b).join(info[["etiqueta", "codigo"]])
            rk = rankings(m, args.top, args.umbral)

            X_t, X_b = x_t.sum(), x_b.sum()
            crec = (X_t - X_b) / abs(X_b) * 100 if X_b else np.nan
            bloques.append(dict(
                hoja=f"{SIGLA[cuenta]}_{lado}", rankings=rk, top=args.top, umbral=args.umbral,
                titulo=f"{nombre} — {nombre_lado}: {t} vs {b}",
                subtitulo=(f"Total {t}: {X_t:,.1f} | Total {b}: {X_b:,.1f} USD millones | "
                           f"Crecimiento: {crec:,.1f}% | Partidas de máximo detalle: {len(m)}"),
            ))
            detalle.append(m.reset_index().assign(cuenta=nombre, lado=nombre_lado))

            if lado == "neto":
                oficial_t = ancho.loc[orden_total, qt].sum()
                oficial_b = ancho.loc[orden_total, qb].sum()
                control.append(dict(cuenta=nombre, periodo=t, total_banrep=oficial_t,
                                    suma_hojas=X_t, diferencia=oficial_t - X_t))
                control.append(dict(cuenta=nombre, periodo=b, total_banrep=oficial_b,
                                    suma_hojas=X_b, diferencia=oficial_b - X_b))

            # consola
            print(f"\n=== {nombre} — {nombre_lado} | {t} vs {b} | total {X_t:,.0f} ({crec:+.1f}%) ===")
            for clave, tabla in rk.items():
                print(f"  {TITULOS[clave].format(n=args.top, u=args.umbral)}")
                for _, f in tabla.iterrows():
                    print(f"    {f.etiqueta[:70]:70s} part {f.participacion_t:7.1f}%  "
                          f"crec {f.crecimiento:8.1f}%  contrib {f.contribucion_pp:7.1f} pp")

    det = pd.concat(detalle, ignore_index=True).rename(columns=encabezados(t, b))
    det = det[["cuenta", "lado", "Orden", "Partida", "Código BanRep"] +
              [c for c in det.columns if c not in ("cuenta", "lado", "Orden", "Partida", "Código BanRep")]]
    ctrl = pd.DataFrame(control)
    notas = [
        f"Fuente: Banco de la República, Balanza de pagos — máximo nivel de detalle "
        f"(data/raw/banrep). Estado de la información en {t}: {', '.join(estados)}.",
        f"Periodo analizado: {t} ({', '.join(qt)}); base: {b} ({', '.join(qb)}). "
        f"Valores en USD millones, suma de los trimestres.",
        "Partidas = máximo nivel de detalle consistente: se baja en la jerarquía solo mientras "
        "los componentes suman el agregado; se excluyen memorandos ('De las cuales', partidas "
        "informativas, autoridades monetarias según corresponda) y desgloses duplicados.",
        "Cuenta corriente: neto = crédito − débito. Cuenta financiera: neto = activos − pasivos "
        "(en el neto los pasivos entran con signo negativo).",
        "Participación (%) = x_t / X_t. Con signos mixtos puede ser negativa o superar 100%; por eso "
        "el ranking de participación ordena por participación absoluta |x_t| / Σ|x_t|.",
        "Crecimiento (%) = (x_t − x_b) / |x_b|. Con saldos negativos, crecimiento positivo = "
        "el saldo sube (p. ej. menor déficit / menor salida neta).",
        f"Ranking de crecimiento: solo partidas con participación absoluta ≥ {args.umbral}% en t o en b, "
        "para evitar crecimientos enormes de partidas con base casi nula.",
        "Contribución (p.p.) = (x_t − x_b) / |X_b|; la suma de contribuciones es el crecimiento del total. "
        "Ranking por valor absoluto (el signo indica si empuja hacia arriba o hacia abajo).",
        "Hoja 'Control': total oficial BanRep vs suma de partidas (debe ser ≈ 0).",
    ]

    OUTPUT.mkdir(exist_ok=True)
    ruta = OUTPUT / f"top_partidas_bp_{t}_vs_{b}.xlsx"
    escribir_excel(ruta, bloques, det, ctrl, notas, t, b)
    print(f"\nControl (total BanRep − suma de partidas):\n{ctrl.round(2).to_string(index=False)}")
    print(f"\nArchivo: {ruta.relative_to(ROOT)}")
    return ruta


if __name__ == "__main__":
    main()
